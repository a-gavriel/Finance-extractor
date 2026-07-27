from openpyxl import load_workbook, Workbook
import tkinter as tk
from tkinter import filedialog
from datetime import datetime
from lib.categorizer import categorize_transaction

# Hide the main Tkinter window
root = tk.Tk()
root.withdraw()

# Open the file dialog and capture the selected file path
file_path = filedialog.askopenfilename(
    title="Select a File",
    initialdir="/",
    filetypes=(("Excel Files", "*.xlsx"), ("All Files", "*.*"))
)

INPUT_FILE = file_path
OUTPUT_FILE = file_path.replace(".xlsx", "_processed.xlsx")


def load_input(path):
    workbook = load_workbook(path)
    return workbook.active


def create_output():
    workbook = Workbook()
    worksheet = workbook.active
    return workbook, worksheet


def write_headers(ws):
    ws.append([
        "Date",
        "Description",
        "Category",
        "Price Cobrado",
        "Price USD",
        "Price CRC",
        "Bank",
        "Card",
    ])


def normalize_transactions(source_ws, target_ws):

    tarjeta_actual = None

    for row in source_ws.iter_rows(min_row=2, values_only=True):

        if not any(row):
            continue

        col1 = str(row[0]).strip() if row[0] else ""

        # Cambio de tarjeta
        if col1.lower() == "tarjeta número:":
            tarjeta_actual = str(row[1]).strip()[-4:]
            continue

        _, fecha, descripcion, monto, moneda, tipo = row[:6]

        # Eliminar créditos
        if tipo == "CREDITO":
            continue

        # Fecha
        fecha = datetime.strptime(fecha, "%d/%m/%Y").strftime("%Y-%m-%d")

        categoria = categorize_transaction(descripcion)

        # Monto
        monto = int(float(str(monto).replace(",", "")))

        monto_crc = float(monto) if moneda == "CRC" else 0
        monto_usd = float(monto) if moneda == "USD" else 0
        bank = "Davibank"

        monto_cobrado = f"{moneda} {monto}"

        target_ws.append([
            fecha,
            descripcion,
            categoria,
            monto_cobrado,
            monto_usd,
            monto_crc,
            bank,
            tarjeta_actual
        ])


def save_output(workbook, path):
    workbook.save(path)


def main():

    source_ws = load_input(INPUT_FILE)

    output_wb, output_ws = create_output()

    write_headers(output_ws)

    normalize_transactions(source_ws, output_ws)

    save_output(output_wb, OUTPUT_FILE)

    print(f"Archivo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()